import datetime
import itertools
import tempfile
import tomllib
import unittest
from dataclasses import replace
from pathlib import Path

from class_schedule import LectureLabClass, NormalClass, Schedule, Section, teaching_loads
from class_schedule.config_inference import infer_configuration_from_template
from class_schedule.config_schema import CourseRelationshipSchema, CoursesFileSchema
from class_schedule.overrides import OverrideEdit, OverrideFile, apply_overrides
from class_schedule.schedule_io import read_schedule
from class_schedule.schedule_model import GroupingError, check_atomic_class_rules, check_conflicts
from class_schedule.solver import SolverConfig, MeetingPattern, RoomRecord, InfeasibleSchedule, solve
from class_schedule.solver.candidates import section_candidates
from class_schedule.solver.constraints import build_slots
from class_schedule.webapp import _serialize_schedule


def rows():
    base = dict(subject="CHEM", number="3245", section="1", instructor="Alice", building="Science")
    return (
        Section(**base, type="CLAS", time_slot="MWF 9:00am", duration=50, room="10"),
        Section(**base, type="LAB", time_slot="M 1:00pm", duration=50, room="20"),
        Section(**base, type="LAB", time_slot="M 1:00pm", duration=170, room="30"),
    )


def config():
    return SolverConfig(persons={}, preferences={}, rooms=[RoomRecord("Science", "99")],
        meeting_patterns=[
            MeetingPattern("MWF", 50, (datetime.time(9), datetime.time(10)), frozenset({"lecture_lab_lecture"})),
            MeetingPattern("M", 170, (datetime.time(13), datetime.time(16)), frozenset({"lecture_lab_lab"})),
        ])


class LectureLabTests(unittest.TestCase):
    def test_identifies_all_source_orders_and_preserves_three_rows(self):
        for source in itertools.permutations(rows()):
            records = [s.to_record() for s in source]
            for record in records:
                record["Meeting Type"] = record.pop("Type")
            item, = Schedule.from_records(records).classes
            self.assertIsInstance(item, LectureLabClass)
            self.assertEqual([s.room for s in item.sections], [s.room for s in source])
            self.assertEqual([s.duration for s in item.scheduling_entries()], [50, 170])
            self.assertEqual(len(item.to_records()), 3)
            self.assertEqual(teaching_loads(Schedule([item])), {"Alice": 5})

    def test_rejects_near_misses(self):
        for changes in (
            {"duration": 250}, {"time_slot": "T 1:00pm"}, {"time_slot": "M 2:00pm"},
            {"room": "20"}, {"instructor": "Bob"}, {"type": "CLAS"},
        ):
            lecture, short, long = rows()
            with self.subTest(changes=changes), self.assertRaises(GroupingError):
                Schedule.from_records([s.to_record() for s in (lecture, short, replace(long, **changes))])

    def test_instructor_edit_links_all_rows_through_both_interfaces(self):
        item = LectureLabClass(rows())
        for index in range(3):
            for updated in (item.change_instructor("Bob", record=index), item.apply_edit("instructor", index, instructor="Bob")):
                self.assertEqual({s.instructor for s in updated.sections}, {"Bob"})
        self.assertEqual({s.instructor for s in item.sections}, {"Alice"})

    def test_fixed_fields_reject_changes_and_do_not_rebase_on_replace(self):
        item = LectureLabClass(rows())
        for index in range(3):
            with self.assertRaisesRegex(ValueError, "fixed"):
                item.change_room("99", record=index)
            with self.assertRaisesRegex(ValueError, "fixed"):
                item.apply_edit("room", index, room="99", building="Other")
        for index in (1, 2):
            with self.assertRaisesRegex(ValueError, "fixed"):
                item.change_time("M 4:00pm", record=index)
        with self.assertRaisesRegex(ValueError, "fixed"):
            replace(item, sections=(replace(item.lecture, room="99"), *item.sections[1:]))
        moved = item.change_time("MWF 10:00am", record=0)
        self.assertEqual(moved.lecture.time_slot, "MWF 10:00am")
        self.assertEqual(moved.lab_long.time_slot, "M 1:00pm")

    def test_baseline_survives_serialization_and_cleaning(self):
        import pandas as pd
        from class_schedule.data_cleaning import clean_dataframe
        item = LectureLabClass(rows())
        data = item.to_records()
        cleaned = clean_dataframe(pd.DataFrame(data))
        self.assertFalse(cleaned.warnings)
        restored, = Schedule.from_dataframe(cleaned.normalized).classes
        self.assertEqual(restored.fixed_locations, item.fixed_locations)
        data[0]["Room"] = "99"
        with self.assertRaisesRegex(GroupingError, "fixed"):
            Schedule.from_records(data)
        data = item.to_records()
        for record in data[1:]:
            record["Time Slot"] = "M 4:00pm"
        with self.assertRaisesRegex(GroupingError, "fixed"):
            Schedule.from_records(data)

    def test_future_lab_unlock_moves_both_rows_and_preserves_durations(self):
        item = LectureLabClass(rows(), lab_time_editable=True)
        moved = item.apply_edit("time", 1, time_slot="M 4:00pm", duration=170)
        self.assertEqual([s.duration for s in moved.sections], [50, 50, 170])
        self.assertEqual(moved.lab_short.time_slot, moved.lab_long.time_slot)
        relation = CourseRelationshipSchema(kind="lecture_lab", members=["CHEM 3245 1"], lab_time_editable=True)
        restored, = Schedule.from_records(moved.to_records(), relationships=[relation]).classes
        self.assertTrue(restored.lab_time_editable)
        self.assertEqual(restored.fixed_lab_slot, "M 1:00pm")

    def test_internal_lecture_lab_overlap_is_reported(self):
        item = LectureLabClass(rows()).change_time("M 1:00pm", record=0)
        self.assertFalse(item.validate())
        self.assertEqual(check_atomic_class_rules(Schedule([item]))[0].rule, "lecture_lab_invalid")

    def test_candidates_keep_rooms_and_lab_time_but_can_change_teacher(self):
        from class_schedule.schedule_model import PersonRecord
        item = LectureLabClass(rows())
        cfg = config()
        cfg = replace(cfg, persons={"Bob": PersonRecord("Bob", 5, ("CHEM 3245",))})
        candidates = section_candidates(item, item.lab_long, cfg, 10)
        self.assertEqual({c.instructor for c in candidates}, {"Alice", "Bob"})
        self.assertEqual({(c.room, c.time_slot, c.duration) for c in candidates}, {("30", "M 1:00pm", 170)})

    def test_two_decisions_generate_three_room_occupancies(self):
        item = LectureLabClass(rows())
        entries = list(item.scheduling_entries())
        candidates = [section_candidates(item, s, config(), 10) for s in entries]
        slots = build_slots(entries, [0, 0], candidates, class_list=[item])
        lab_slots = [s for s in slots if s.section == 1]
        self.assertEqual(len(lab_slots), 2)
        short = next(s for s in lab_slots if s.room_key.endswith("20"))
        long = next(s for s in lab_slots if s.room_key.endswith("30"))
        self.assertEqual(short.end, datetime.time(13, 50))
        self.assertEqual(long.end, datetime.time(15, 50))
        self.assertEqual(short.instructor, "")
        self.assertEqual(long.instructor, "Alice")
        self.assertEqual((short.section, short.candidate), (long.section, long.candidate))

    def test_short_room_hard_rules_and_preferences_are_not_lost(self):
        from class_schedule.schedule_model import ConstraintRule, PreferenceRule, check_soft_preferences
        item = LectureLabClass(rows())
        cfg = replace(config(), constraint_rules=(ConstraintRule(direction="-", course="CHEM 3245", room="Science 20"),))
        self.assertEqual(section_candidates(item, item.lab_long, cfg, 10), [])
        rules = (
            PreferenceRule(course="CHEM 3245", direction="dislike", weight=7),
            PreferenceRule(room="Science 20", direction="dislike", weight=11),
        )
        cfg = replace(config(), global_rules=rules)
        candidates = [section_candidates(item, s, cfg, 10) for s in item.scheduling_entries()]
        # General rule is charged twice (lecture + long lab), short room once.
        self.assertEqual(sum(min(c.cost for c in values) for values in candidates), 25)
        _, findings = check_soft_preferences(Schedule([item]), {}, {}, global_rules=rules)
        self.assertEqual(sum(f.penalty for f in findings if f.rule == "custom_rule"), 25)

    def test_short_row_lock_constrains_long_lab_decision_in_shuffled_order(self):
        item = LectureLabClass(tuple(reversed(rows())), lab_time_editable=True)
        blocker = NormalClass((replace(item.lecture, number="1003", instructor="Bob", room="20", time_slot="M 1:30pm"),))
        cfg = config()
        cfg.meeting_patterns.append(MeetingPattern("M", 50, (datetime.time(13, 30),), frozenset({"normal"})))
        locks = {
            (item.course_ids[0], 1): frozenset({"time"}),
            (blocker.course_ids[0], None): frozenset({"time", "room", "building"}),
        }
        with self.assertRaises(InfeasibleSchedule):
            solve(Schedule([item, blocker]), cfg, locks=locks, search_workers=1, time_limit_seconds=5)

    def test_solver_preserves_type_baseline_and_three_rows(self):
        item = LectureLabClass(rows())
        solved = solve(Schedule([item]), config(), search_workers=1, time_limit_seconds=5)
        result, = solved.classes
        self.assertIsInstance(result, LectureLabClass)
        self.assertEqual(len(result.sections), 3)
        self.assertEqual(result.fixed_locations, item.fixed_locations)
        self.assertEqual(result.lab_long.time_slot, item.lab_long.time_slot)
        self.assertFalse(check_conflicts(solved))

    def test_short_room_blocks_only_first_fifty_minutes(self):
        item = LectureLabClass(rows())
        cfg = SolverConfig(persons={}, preferences={}, rooms=[], meeting_patterns=[])
        for start, allowed in (("M 1:30pm", False), ("M 1:50pm", True), ("M 2:00pm", True)):
            other = NormalClass((replace(item.lecture, number="1003", instructor="Bob", room="20", time_slot=start),))
            schedule = Schedule([item, other])
            with self.subTest(start=start):
                if allowed:
                    solved = solve(schedule, cfg, search_workers=1, time_limit_seconds=5)
                    self.assertFalse(check_conflicts(solved))
                else:
                    with self.assertRaises(InfeasibleSchedule):
                        solve(schedule, cfg, search_workers=1, time_limit_seconds=5)

    def test_solver_can_move_lecture_out_of_lab(self):
        item = LectureLabClass(rows()).change_time("M 1:00pm", record=0)
        solved = solve(Schedule([item]), config(), search_workers=1, time_limit_seconds=5)
        self.assertTrue(solved.classes[0].validate())

    def test_unlock_allows_solver_to_move_whole_lab(self):
        item = LectureLabClass(rows(), lab_time_editable=True)
        blocker = NormalClass((replace(item.lecture, number="1003", instructor="Bob", room="20", time_slot="M 1:30pm"),))
        cfg = config()
        cfg.meeting_patterns.append(MeetingPattern("M", 50, (datetime.time(13, 30),), frozenset({"normal"})))
        solved = solve(Schedule([item, blocker]), cfg, locks={(blocker.course_ids[0], None): frozenset({"room", "building", "time"})}, search_workers=1, time_limit_seconds=5)
        result = solved.classes[0]
        self.assertEqual(result.lab_long.time_slot, "M 4:00pm")
        self.assertEqual(result.lab_short.time_slot, "M 4:00pm")
        self.assertTrue(result.lab_time_editable)

    def test_overrides_enforce_locks_and_link_instructor(self):
        schedule = Schedule([LectureLabClass(rows())])
        edited = apply_overrides(schedule, OverrideFile(edits=(OverrideEdit("CHEM 3245-1", record=1, instructor="Bob"),)))
        self.assertEqual({s.instructor for s in edited.classes[0].sections}, {"Bob"})
        with self.assertRaisesRegex(ValueError, "fixed"):
            apply_overrides(schedule, OverrideFile(edits=(OverrideEdit("CHEM 3245-1", record=1, time_slot="M 4:00pm"),)))

    def test_view_metadata_preserves_source_indexes(self):
        item = LectureLabClass(tuple(reversed(rows())))
        payload, = _serialize_schedule(Schedule([item]))
        self.assertEqual(payload["scheduling_records"], [0, 2])
        self.assertEqual(payload["editable_fields"], [["instructor"], ["instructor"], ["instructor", "time"]])
        self.assertTrue(payload["linked_fields"]["instructor"])

    def test_uploaded_template_can_be_edited_and_exported_via_web(self):
        import pandas as pd
        from fastapi.testclient import TestClient
        from unittest.mock import patch
        from class_schedule import webapp
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = pd.DataFrame([s.to_record() for s in rows()]).to_csv(index=False).encode()
            with patch.object(webapp, "CONFIG_DIR", root / "config"), patch.object(webapp, "WORK_ROOT", root / "work"):
                client = TestClient(webapp.create_app())
                response = client.post("/api/configuration-packages", data={"current_package": "TEST"},
                    files={"config_files": ("lecture-lab.csv", source, "text/csv")})
                self.assertEqual(response.status_code, 200, response.text)
                cfg = SolverConfig.load(root / "config", package="TEST")
                current = read_schedule(root / "work/TEST/initial/initial.csv", relationships=cfg.courses.active_relationships)
                self.assertIsInstance(current.classes[0], LectureLabClass)
                records = _serialize_schedule(current)[0]["sections"]

                def request(field, value, record_index):
                    target = records[record_index]
                    return client.post("/api/edit", json={
                        "package": "TEST", "records": records, "class_index": 0, "record_index": record_index,
                        "expected_course_ids": ["CHEM 3245-1"],
                        "expected_record": {"subject": "CHEM", "number": "3245", "section": "1",
                            "expected_time_slot": target["Time Slot"]},
                        "field": field, "value": value,
                    })

                response = request("instructor", "Bob", 1)
                self.assertEqual(response.status_code, 200, response.text)
                self.assertEqual({r["Instructor"] for r in response.json()["classes"][0]["sections"]}, {"Bob"})
                response = request("room", {"building": "Science", "room": "99"}, 0)
                self.assertEqual(response.status_code, 400, response.text)
                response = request("time", {"days": "M", "start": "16:00"}, 1)
                self.assertEqual(response.status_code, 400, response.text)
                response = request("time", {"days": "MWF", "start": "10:00"}, 0)
                self.assertEqual(response.status_code, 200, response.text)
                self.assertEqual(response.json()["classes"][0]["sections"][0]["Duration"], 50)
                instructor = root / "instructor.xlsx"
                room = root / "room.xlsx"
                raw = root / "raw.xlsx"
                current.to_instructor_excel(instructor)
                current.to_room_excel(room)
                current.to_raw_excel(raw)
                instructor_book = load_workbook(instructor)
                text = " ".join(str(c.value) for ws in instructor_book for row in ws for c in row if c.value)
                self.assertNotIn("Science 20", text)
                self.assertIn("Science 20", load_workbook(room).sheetnames)
                self.assertEqual(load_workbook(raw).active.max_row, 4)

    def test_excluding_previous_solution_still_uses_two_decisions(self):
        original = Schedule([LectureLabClass(rows())])
        first = solve(original, config(), search_workers=1, time_limit_seconds=5)
        second = solve(original, config(), previous=first, search_workers=1, time_limit_seconds=5)
        self.assertNotEqual(first.classes[0].lecture.time_slot, second.classes[0].lecture.time_slot)
        self.assertEqual(first.classes[0].lab_long.time_slot, second.classes[0].lab_long.time_slot)

    def test_real_workbook_inference_loads_and_reconciles(self):
        from class_schedule.reconciliation import reconcile_records
        source = Path(__file__).parents[1] / "config/27/202720 maps.xlsx"
        if not source.is_file():
            self.skipTest("Local user workbook is not part of the test fixtures")
        original = read_schedule(source)
        self.assertEqual(sum(isinstance(c, LectureLabClass) for c in original), 3)
        inferred = infer_configuration_from_template(source, package="TEST")
        parsed = CoursesFileSchema.model_validate(tomllib.loads(inferred.files["courses.toml"].decode()))
        self.assertEqual(sum(r.kind == "lecture_lab" for r in parsed.relationships), 3)
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            package = root / "TEST"
            package.mkdir()
            for name, content in inferred.files.items():
                destination = package / ("basicinfo" if name in {"catalogs.toml", "locations.toml", "timeslot.toml", "persons.toml"} else "") / name
                destination.parent.mkdir(exist_ok=True)
                destination.write_bytes(content)
            cfg = SolverConfig.load(root, package="TEST")
            rebuilt, _ = reconcile_records(original.to_records(), cfg)
            self.assertEqual(sum(isinstance(c, LectureLabClass) for c in rebuilt), 3)


if __name__ == "__main__":
    unittest.main()
